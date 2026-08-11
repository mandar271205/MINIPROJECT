"""Step 8a – Output reports.

Generate human-readable attendance reports in XLSX and PDF format.
Uses ``openpyxl`` for Excel and ``reportlab`` for PDF.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .models import AttendanceDecision, AttendanceReport

# -------------------------------------------------------------------- #
#  Helpers                                                               #
# -------------------------------------------------------------------- #

_STATUS_COLORS = {
    "present": "4CAF50",      # green
    "late-absent": "FF9800",  # orange
    "absent": "F44336",       # red
}


def _ts(seconds: float | None) -> str:
    if seconds is None:
        return "--"
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


# -------------------------------------------------------------------- #
#  XLSX generation                                                       #
# -------------------------------------------------------------------- #

def generate_xlsx(
    report: AttendanceReport,
    output_path: str | Path | None = None,
) -> BytesIO | Path:
    """
    Write the attendance report to an Excel workbook.

    If *output_path* is given the file is saved to disk and the
    :class:`Path` is returned.  Otherwise a :class:`BytesIO` buffer
    is returned (useful for streaming downloads).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # ---------- styles ----------
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="C10202", end_color="C10202", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")

    # ---------- title row ----------
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"Attendance Report – Session #{report.session_id}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = center

    ws.merge_cells("A2:E2")
    ws["A2"].value = f"Generated: {report.created_at:%Y-%m-%d %H:%M}"
    ws["A2"].alignment = center

    # ---------- header row ----------
    headers = ["#", "Student ID", "Name", "Status", "First Seen"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center

    # ---------- data rows ----------
    for i, row in enumerate(report.rows, start=1):
        r = i + 4
        ws.cell(row=r, column=1, value=i).alignment = center
        ws.cell(row=r, column=2, value=row.student_id).alignment = center
        ws.cell(row=r, column=3, value=row.student_name)
        status_cell = ws.cell(row=r, column=4, value=row.status.upper())
        status_cell.alignment = center
        color = _STATUS_COLORS.get(row.status, "CCCCCC")
        status_cell.fill = PatternFill(
            start_color=color, end_color=color, fill_type="solid"
        )
        status_cell.font = Font(bold=True, color="FFFFFF")
        ws.cell(row=r, column=5, value=_ts(row.first_seen_sec)).alignment = center

        for col in range(1, 6):
            ws.cell(row=r, column=col).border = thin_border

    # ---------- column widths ----------
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 14

    # ---------- summary row ----------
    summary_row = len(report.rows) + 6
    present = sum(1 for r in report.rows if r.status == "present")
    late = sum(1 for r in report.rows if r.status == "late-absent")
    absent = sum(1 for r in report.rows if r.status == "absent")
    ws.cell(row=summary_row, column=1, value="Summary:").font = Font(bold=True)
    ws.cell(row=summary_row, column=2, value=f"Present: {present}")
    ws.cell(row=summary_row, column=3, value=f"Late-Absent: {late}")
    ws.cell(row=summary_row, column=4, value=f"Absent: {absent}")
    ws.cell(row=summary_row, column=5, value=f"Total: {len(report.rows)}")

    if output_path:
        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(p))
        return p

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# -------------------------------------------------------------------- #
#  PDF generation                                                        #
# -------------------------------------------------------------------- #

def generate_pdf(
    report: AttendanceReport,
    output_path: str | Path | None = None,
) -> BytesIO | Path:
    """
    Write the attendance report to a PDF document.
    """
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf if output_path is None else str(output_path),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=20 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    elements: list = []

    # Title
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=18,
        textColor=colors.HexColor("#C10202"),
        spaceAfter=6,
    )
    elements.append(
        Paragraph(f"Attendance Report – Session #{report.session_id}", title_style)
    )
    elements.append(
        Paragraph(
            f"Generated: {report.created_at:%Y-%m-%d %H:%M}",
            styles["Normal"],
        )
    )
    elements.append(Spacer(1, 12))

    # Table
    header = ["#", "Student ID", "Name", "Status", "First Seen"]
    data = [header]
    for i, row in enumerate(report.rows, start=1):
        data.append(
            [
                str(i),
                str(row.student_id),
                row.student_name,
                row.status.upper(),
                _ts(row.first_seen_sec),
            ]
        )

    table = Table(data, colWidths=[30, 70, 150, 80, 70])
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C10202")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
    ]

    # Color-code status cells
    for i, row in enumerate(report.rows, start=1):
        color = _STATUS_COLORS.get(row.status, "CCCCCC")
        style_cmds.append(
            ("BACKGROUND", (3, i), (3, i), colors.HexColor(f"#{color}"))
        )
        style_cmds.append(("TEXTCOLOR", (3, i), (3, i), colors.white))

    table.setStyle(TableStyle(style_cmds))
    elements.append(table)

    # Summary
    present = sum(1 for r in report.rows if r.status == "present")
    late = sum(1 for r in report.rows if r.status == "late-absent")
    absent = sum(1 for r in report.rows if r.status == "absent")
    elements.append(Spacer(1, 12))
    elements.append(
        Paragraph(
            f"<b>Summary:</b> Present: {present} | "
            f"Late-Absent: {late} | Absent: {absent} | "
            f"Total: {len(report.rows)}",
            styles["Normal"],
        )
    )

    doc.build(elements)

    if output_path:
        return Path(output_path)

    buf.seek(0)
    return buf
